from odoo import api, models, fields, _
from odoo.exceptions import UserError

import base64
import io
from openpyxl import load_workbook


class ImportExcelWizard(models.TransientModel):
    _name = 'import.excel.wizard'
    _description = 'Import Excel Wizard'

    file = fields.Binary("File Excel", required=True)
    filename = fields.Char("File Name")
    line_ids = fields.One2many('import.excel.line', 'wizard_id', string="Preview Data")
    exam_id = fields.Many2one('student.exam', string="Exam")

    @api.model
    def default_get(self, fields_list):
        res = super().default_get(fields_list)
        # Lưu exam_id từ context vào field để không bị mất
        if self.env.context.get('record_id'):
            res['exam_id'] = self.env.context['record_id']
        return res

    # ====================== LOAD FILE ======================
    def action_load_file(self):
        self.line_ids.unlink()

        wb = load_workbook(filename=io.BytesIO(base64.b64decode(self.file)))
        sheet = wb.active
        lines = []

        for row in sheet.iter_rows(min_row=5, values_only=True):
            if not row or all(cell is None for cell in row[:2]):
                continue
            name = str(row[0]).strip() if row[0] else False
            score = row[1] if len(row) > 1 else False

            error = "Missing name" if not name else ""
            try:
                score_val = float(score) if score not in [None, False, "", " "] else 0.0
            except Exception:
                score_val = 0.0
                error = "Invalid score"

            lines.append((0, 0, {'name': name, 'score': score_val, 'error': error}))

        self.line_ids = lines

        # Giữ nguyên context khi mở lại wizard
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
            'context': self.env.context,
        }

    # ====================== IMPORT ======================
    def action_import(self):
        exam = self.exam_id
        if not exam:
            raise UserError(_("Không tìm thấy Exam! Hãy mở wizard từ form Exam."))

        if not exam.exists():
            raise UserError(_("Bản ghi Exam không tồn tại hoặc đã bị xóa."))

        created_count = 0
        for line in self.line_ids:
            if line.error:
                continue
            self.env['student.exam.line'].create({
                'exam_id': exam.id,
                'subject': line.name,
                'score': line.score or 0.0,
            })
            created_count += 1

        # Quay về form Exam để xem kết quả ngay
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'student.exam',
            'view_mode': 'form',
            'res_id': exam.id,
            'target': 'current',
        }


class ImportExcelLine(models.TransientModel):
    _name = 'import.excel.line'
    _description = 'Import Excel Line'

    wizard_id = fields.Many2one('import.excel.wizard')
    name = fields.Char("Subject")
    score = fields.Float("Score")
    error = fields.Char("Error")